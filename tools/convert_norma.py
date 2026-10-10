#!/usr/bin/env python3
"""Konversi Tabel_Nilai_Jasmani_Polri_2026-v2.xlsx -> norma_polri_2026_v2.json

Prinsip:
- Status per baris dipertahankan apa adanya (VERIFIED / LINEAR_ESTIMATE). Tidak ada yang dinaikkan menjadi resmi.
- Metadata yang tidak ada di workbook diisi null, bukan ditebak.
- Arah penilaian dideklarasikan eksplisit di sini lalu DIVERIFIKASI terhadap data; jika tidak cocok, skrip gagal.

Pemakaian:
  python3 tools/convert_norma.py source/Tabel_Nilai_Jasmani_Polri_2026-v2.xlsx \
      domain/src/main/resources/norms/norma_polri_2026_v2.json
"""
import hashlib
import json
import re
import sys

from openpyxl import load_workbook

HEADERS = {  # nama kolom -> (id, kelompok, arah)
    "Lari 12M": ("LARI_12M", "A", "HIGHER_IS_BETTER"),
    "Pull Up": ("PULL_UP", "B", "HIGHER_IS_BETTER"),
    "Chinning": ("CHINNING", "B", "HIGHER_IS_BETTER"),
    "Sit Up": ("SIT_UP", "B", "HIGHER_IS_BETTER"),
    "Push Up": ("PUSH_UP", "B", "HIGHER_IS_BETTER"),
    "Shuttle Run": ("SHUTTLE_RUN", "B", "LOWER_IS_BETTER"),
    "Renang": ("RENANG", "C", "LOWER_IS_BETTER"),
}
SHEETS = {
    "Norma Pria": ("pria_2026_v2", "PRIA", "Pria"),
    "Norma Wanita": ("wanita_2026_v2", "WANITA", "Wanita"),
}
CHINNING_NOTE = (
    "Satuan repetisi dikonfirmasi pengguna. Namun 72 repetisi pada nilai 100 tampak tidak lazim; "
    "belum diverifikasi terhadap dokumen resmi."
)


def parse_header(h):
    m = re.fullmatch(r"(.+?) \((\w+)\)", h.strip())
    if not m or m.group(1) not in HEADERS:
        raise SystemExit(f"Header tidak dikenali: {h!r}")
    return m.group(1), m.group(2)


def parse_status(s):
    if s == "Estimasi Linier":
        return "LINEAR_ESTIMATE", None
    m = re.fullmatch(r"Terverifikasi \((.+)\)", s or "")
    if m:
        return "VERIFIED", m.group(1)
    raise SystemExit(f"Status tidak dikenali: {s!r}")


def main(src, dst):
    sha = hashlib.sha256(open(src, "rb").read()).hexdigest()
    wb = load_workbook(src, data_only=True)
    audit = {"sheetsRead": len(wb.sheetnames), "sheetNames": wb.sheetnames, "tables": {}, "warnings": []}
    tables = []
    for sheet, (tid, gender, cat) in SHEETS.items():
        ws = wb[sheet]
        rows_raw = [[c.value for c in r] for r in ws.iter_rows(min_row=1)]
        hdr, body = rows_raw[0], rows_raw[1:]
        if hdr[0] != "Nilai" or hdr[-1] != "Status":
            raise SystemExit(f"Kolom tidak cocok di {sheet}")
        comps = []
        for h in hdr[1:-1]:
            name, unit = parse_header(h)
            cid, grp, direction = HEADERS[name]
            comps.append({
                "id": cid, "header": h, "unit": unit, "direction": direction, "group": grp,
                "unitNote": CHINNING_NOTE if cid == "CHINNING" else None,
            })
        rows, seen, empty = [], set(), 0
        for r in body:
            empty += sum(v in (None, "") for v in r)
            score = int(r[0])
            if score in seen:
                raise SystemExit(f"Nilai duplikat {score} di {sheet}")
            seen.add(score)
            status, ref = parse_status(r[-1])
            rows.append({
                "score": score, "status": status, "reference": ref,
                "values": {c["id"]: float(v) for c, v in zip(comps, r[1:-1])},
            })
        rows.sort(key=lambda x: -x["score"])
        for c in comps:  # verifikasi arah terhadap data
            col = [x["values"][c["id"]] for x in rows]
            if c["direction"] == "HIGHER_IS_BETTER":
                ok = all(a >= b for a, b in zip(col, col[1:]))
            else:
                ok = all(a <= b for a, b in zip(col, col[1:]))
            if not ok:
                raise SystemExit(f"Arah {c['id']} di {sheet} tidak cocok dengan data")
        ver = [x["score"] for x in rows if x["status"] == "VERIFIED"]
        est = len(rows) - len(ver)
        if sorted(r["score"] for r in rows) != list(range(101)):
            audit["warnings"].append(f"{sheet}: nilai tidak lengkap 0..100")
        note = (
            f"{len(ver)} titik terverifikasi (nilai {', '.join(map(str, ver))}) dengan rujukan Kep/1352/VI/2020; "
            f"{est} baris estimasi linier. Status di workbook tidak membuktikan norma berlaku untuk seluruh "
            f"seleksi Akpol 2026/2027."
        )
        if gender == "WANITA":
            audit["warnings"].append(
                "Norma Wanita: Chinning bernilai 0 pada nilai 0-40; satuan perlu verifikasi dokumen resmi.")
        audit["tables"][tid] = {
            "rows": len(rows), "emptyCells": empty, "duplicates": 0,
            "verifiedScores": ver, "estimateRows": est,
        }
        tables.append({
            "id": tid, "gender": gender, "sheet": sheet,
            "metadata": {
                "sourceName": "Tabel_Nilai_Jasmani_Polri_2026-v2.xlsx (file impor pengguna)",
                "documentName": None, "decisionNumber": None, "documentDate": None,
                "validPeriod": None, "category": cat, "applicability": "UNCONFIRMED",
                "verificationNote": note,
            },
            "components": comps, "rows": rows,
        })
    log = []
    for r in wb["Log Pengujian Batas"].iter_rows(min_row=2, values_only=True):
        log.append({"id": r[0], "gender": r[1], "component": r[2], "input": r[3],
                    "score": r[4], "change": r[5], "evidence": r[6]})
    audit["boundaryLogRows"] = len(log)
    audit["warnings"].append(
        "Log Pengujian Batas: sumber bukti ('Blangko TKJ Polri', 'Standar SSDM Polri') berbeda dari rujukan "
        "di sheet norma (Kep/1352/VI/2020); entri Wanita-Renang tidak ada.")
    out = {
        "schemaVersion": 1,
        "source": {"fileName": src.split("/")[-1], "sha256": sha, "importedBy": "tools/convert_norma.py"},
        "tables": tables, "boundaryLog": log, "importAudit": audit,
    }
    with open(dst, "w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False, indent=1)
    print(json.dumps(audit, ensure_ascii=False, indent=1))


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
