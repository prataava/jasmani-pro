#!/usr/bin/env python3
"""Implementasi REFERENSI independen kalkulator (membaca langsung dari .xlsx, bukan dari JSON hasil konversi).

Fungsi:
1. Menghasilkan domain/src/test/resources/golden_vectors.json: untuk tiap ambang unik, hasil tepat di ambang,
   sedikit di bawah, dan sedikit di atas, lengkap dengan nilai harapan untuk kedua aturan tie-break.
2. Mencetak angka harapan untuk tes tulisan tangan di Kotlin.

Kode Kotlin TIDAK dibuat dari skrip ini; keduanya ditulis terpisah lalu dibandingkan lewat vektor uji.
Pemakaian: python3 tools/reference_check.py source/Tabel_...xlsx domain/src/test/resources/golden_vectors.json
"""
import json
import sys

from openpyxl import load_workbook

COLS = {  # indeks kolom di sheet -> (id, arah)
    1: ("LARI_12M", "H"), 2: ("PULL_UP", "H"), 3: ("SIT_UP", "H"),
    4: ("PUSH_UP", "H"), 5: ("SHUTTLE_RUN", "L"), 6: ("RENANG", "L"),
}
DELTA = {"LARI_12M": 1.0, "PULL_UP": 1.0, "CHINNING": 1.0, "SIT_UP": 1.0, "PUSH_UP": 1.0,
         "SHUTTLE_RUN": 0.05, "RENANG": 0.05}


def load(path):
    wb = load_workbook(path, data_only=True)
    out = {}
    for sheet, gender in (("Norma Pria", "PRIA"), ("Norma Wanita", "WANITA")):
        rows = [list(r) for r in wb[sheet].iter_rows(min_row=2, values_only=True)]
        cols = dict(COLS)
        if gender == "WANITA":
            cols[2] = ("CHINNING", "H")
        out[gender] = (cols, sorted(rows, key=lambda r: -r[0]))
    return out


def score(cols, rows, idx, raw, tie):
    cid, d = cols[idx]
    ok = (lambda t: raw >= t) if d == "H" else (lambda t: raw <= t)
    match = next((r for r in rows if ok(r[idx])), None)
    if match is None:
        return 0, None, None, "BELOW_TABLE_MIN"
    thr = match[idx]
    tied = [r for r in rows if r[idx] == thr]
    chosen = match if tie == "HIGHEST" else min(tied, key=lambda r: r[0])
    rng = (min(r[0] for r in tied), max(r[0] for r in tied)) if len(tied) > 1 else None
    top = rows[0]
    beyond = match is top and ((raw > thr) if d == "H" else (raw < thr))
    status = "VERIFIED" if str(chosen[7]).startswith("Terverifikasi") else "LINEAR_ESTIMATE"
    return chosen[0], status, rng, "ABOVE_TABLE_MAX" if beyond else "NONE"


def main(src, dst):
    data = load(src)
    vectors = []
    for gender, (cols, rows) in data.items():
        for idx, (cid, _) in cols.items():
            for thr in sorted({r[idx] for r in rows}):
                for raw in {round(thr, 4), round(thr - DELTA[cid], 4), round(thr + DELTA[cid], 4)}:
                    if raw < 0:
                        continue
                    sh, st_h, rng, flag = score(cols, rows, idx, raw, "HIGHEST")
                    sl, st_l, _, _ = score(cols, rows, idx, raw, "LOWEST")
                    vectors.append({
                        "gender": gender, "component": cid, "input": raw,
                        "scoreHighest": sh, "statusHighest": st_h,
                        "scoreLowest": sl, "statusLowest": st_l,
                        "rangeMin": rng[0] if rng else None, "rangeMax": rng[1] if rng else None,
                        "flag": flag,
                    })
    with open(dst, "w", encoding="utf-8") as f:
        json.dump(vectors, f, ensure_ascii=False)
    print("Vektor emas:", len(vectors))

    # ---- angka harapan untuk tes tulisan tangan ----
    pc, pr = data["PRIA"]
    wc, wr = data["WANITA"]

    def s(g, idx, raw, tie="HIGHEST"):
        cols, rows = (pc, pr) if g == "P" else (wc, wr)
        return score(cols, rows, idx, raw, tie)

    print("Pria lari 3444:", s("P", 1, 3444), "| 3445:", s("P", 1, 3445), "| 1327:", s("P", 1, 1327), "| 1326:", s("P", 1, 1326))
    print("Pria pull-up 17 H/L:", s("P", 2, 17), s("P", 2, 17, "LOWEST"))
    print("Pria shuttle 16.2:", s("P", 5, 16.2), "| 16.21:", s("P", 5, 16.21))
    print("Wanita push-up 0 H/L:", s("W", 4, 0), s("W", 4, 0, "LOWEST"))
    print("Wanita chinning 0 H/L:", s("W", 2, 0), s("W", 2, 0, "LOWEST"))
    # contoh lengkap pria
    ex = {"lari": 2809, "pull": 12, "sit": 28, "push": 31, "shuttle": 17.9, "renang": 35.0}
    sc = {k: s("P", i, v)[0] for k, (i, v) in
          {"lari": (1, 2809), "pull": (2, 12), "sit": (3, 28), "push": (4, 31), "shuttle": (5, 17.9), "renang": (6, 35.0)}.items()}
    avg_b = (sc["pull"] + sc["sit"] + sc["push"] + sc["shuttle"]) / 4
    print("Contoh pria", ex, "->", sc, "| rata-rata B:", avg_b, "| gabungan AB:", (sc["lari"] + avg_b) / 2)
    ex2 = {"lari": 3000, "pull": 13, "sit": 31, "push": 33, "shuttle": 17.45, "renang": 30.0}
    sc2 = {k: s("P", i, v)[0] for k, (i, v) in
           {"lari": (1, 3000), "pull": (2, 13), "sit": (3, 31), "push": (4, 33), "shuttle": (5, 17.45), "renang": (6, 30.0)}.items()}
    avg_b2 = (sc2["pull"] + sc2["sit"] + sc2["push"] + sc2["shuttle"]) / 4
    print("Contoh pria #2", ex2, "->", sc2, "| rata-rata B:", avg_b2, "| gabungan AB:", (sc2["lari"] + avg_b2) / 2)
    print("Ambang pria nilai 70 (lari, pull, sit, push, shuttle, renang):", [r for r in pr if r[0] == 70][0][1:7])


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
